# -*- coding: utf-8 -*-
"""Genera sistema/js/datos-mg.js desde los Excel de Andrés (Mantención general)."""
from pathlib import Path
from datetime import datetime, date
import json
import openpyxl

BASE = Path(r"C:\Users\lolma\OneDrive\Escritorio\WORK\CAPSTONE\Enviado por andres\Mantecion General")
OUT = Path(r"C:\Users\lolma\OneDrive\Escritorio\WORK\CAPSTONE\sistema\js\datos-mg.js")


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d-%m-%Y")
    if isinstance(v, date):
        return v.strftime("%d-%m-%Y")
    s = str(v).strip()
    if s in {"0", "None"}:
        return ""
    return " ".join(s.replace("\n", " ").split())


def last_date_in_row(ws, r, c0, c1):
    best = None
    best_c = None
    for c in range(c0, c1 + 1):
        v = ws.cell(r, c).value
        if isinstance(v, (datetime, date)):
            if best is None or v > best:
                best, best_c = v, c
    return fmt(best), best_c


def dump():
    data = {}

    # FOSAS
    wb = openpyxl.load_workbook(BASE / " LISTA FOSAS DE ALCANTARILLADOS .xlsx", data_only=True)
    ws = wb["Hoja1"]
    fosas = []
    for r in range(3, 31):
        nom = fmt(ws.cell(r, 2).value)
        ubi = fmt(ws.cell(r, 3).value)
        if not nom:
            continue
        ultima, c = last_date_in_row(ws, r, 4, 17)
        fosas.append({"punto": nom, "ubicacion": ubi or "—", "ultima": ultima, "tipo": "Grasera" if "grasera" in nom.lower() else ("Planta" if "planta" in nom.lower() else ("Cámara" if "camara" in nom.lower() or "cámara" in nom.lower() else "Fosa"))})
    data["fosas"] = fosas
    data["fosas_contactos"] = [
        {"quien": "Villareal — Laura Guardia", "dato": "+56 9 9817 6284"},
        {"quien": "Bernardo Carrillo (Marcela)", "dato": "9 9441 2474 / 2 2824 2062 — camión 10.000 L Paine"},
        {"quien": "Richard Delard", "dato": "+56 9 8198 8248 — camión 5.000 L"},
    ]
    wb.close()

    # BODEGAS
    wb = openpyxl.load_workbook(BASE / "LISTA DE BODEGAS.xlsx", data_only=True)
    ws = wb["Hoja1"]
    bodegas = []
    for r in range(2, 31):
        nom = fmt(ws.cell(r, 3).value)
        if not nom:
            continue
        bodegas.append({
            "cod": fmt(ws.cell(r, 2).value),
            "nombre": nom,
            "ubicacion": fmt(ws.cell(r, 4).value),
            "ultima": fmt(ws.cell(r, 5).value),
            "quien": fmt(ws.cell(r, 6).value),
        })
    data["bodegas"] = bodegas
    wb.close()

    # MAQUINAS
    wb = openpyxl.load_workbook(BASE / "LISTA DE MANTENCIONES.xlsx", data_only=True)
    ws = wb["MAQUINAS"]
    maq = []
    for r in range(3, 51):
        obj = fmt(ws.cell(r, 2).value)
        if not obj:
            continue
        ultima, c = last_date_in_row(ws, r, 5, 43)
        quien = fmt(ws.cell(r, c + 1).value) if c else ""
        det = fmt(ws.cell(r, c + 2).value) if c else ""
        if quien and len(quien) > 40 and not det:
            det, quien = quien, det
        maq.append({
            "num": fmt(ws.cell(r, 1).value),
            "equipo": obj,
            "ubicacion": fmt(ws.cell(r, 3).value),
            "proveedor": fmt(ws.cell(r, 4).value),
            "ultima": ultima,
            "quien": quien[:80],
            "detalle": det[:140],
        })
    data["maquinas"] = maq

    rad = []
    ws = wb["RADIADORES"]
    for r in range(3, 35):
        hab = fmt(ws.cell(r, 2).value)
        if not hab:
            continue
        rad.append({
            "zona": hab,
            "ubicacion": fmt(ws.cell(r, 3).value),
            "proveedor": fmt(ws.cell(r, 4).value),
            "ultima": fmt(ws.cell(r, 5).value),
            "quien": fmt(ws.cell(r, 6).value),
            "detalle": fmt(ws.cell(r, 7).value),
        })
    data["radiadores"] = rad

    refri = []
    ws = wb["REFRIGERADOR"]
    for r in range(3, 22):
        nom = fmt(ws.cell(r, 2).value)
        if not nom:
            continue
        refri.append({
            "equipo": nom,
            "ubicacion": fmt(ws.cell(r, 3).value),
            "proveedor": fmt(ws.cell(r, 4).value),
            "ultima": fmt(ws.cell(r, 5).value),
            "quien": fmt(ws.cell(r, 6).value),
            "detalle": fmt(ws.cell(r, 7).value),
        })
    data["refrigeradores"] = refri
    wb.close()

    # PUERTAS
    wb = openpyxl.load_workbook(BASE / "LISTA DE PUERTAS POR ENGRASAR Y ACEITAR.xlsx", data_only=True)
    pt = []
    seen = set()
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for r in range(2, (ws.max_row or 1) + 1):
            nom = fmt(ws.cell(r, 2).value)
            if not nom or nom in seen:
                continue
            seen.add(nom)
            pt.append({
                "puerta": nom,
                "ubicacion": fmt(ws.cell(r, 3).value),
                "ultima": fmt(ws.cell(r, 4).value),
                "quien": fmt(ws.cell(r, 6).value),
            })
    data["puertas"] = pt
    wb.close()

    # BANOS
    wb = openpyxl.load_workbook(BASE / "MANTENCION DE BAÑOS.xlsx", data_only=True)
    ws = wb.active
    banos = []
    for r in range(3, 35):
        hab = fmt(ws.cell(r, 1).value)
        if not hab:
            continue
        banos.append({
            "recinto": hab,
            "lavamanos": fmt(ws.cell(r, 2).value),
            "tina": fmt(ws.cell(r, 3).value),
            "porta_rollo": fmt(ws.cell(r, 4).value),
        })
    data["banos"] = banos
    wb.close()

    # PILAS
    wb = openpyxl.load_workbook(BASE / "MANTENCION DE PILAS.xlsx", data_only=True)
    ws = wb.active
    pilas = []
    for r in range(3, 31):
        hab = fmt(ws.cell(r, 1).value)
        if not hab:
            continue
        pilas.append({
            "zona": hab,
            "tv": fmt(ws.cell(r, 2).value) or "Sin fecha",
            "fecha_tv": fmt(ws.cell(r, 3).value),
            "deco": fmt(ws.cell(r, 4).value) or "Sin fecha",
            "fecha_deco": fmt(ws.cell(r, 5).value),
        })
    data["pilas"] = pilas
    wb.close()

    # TV
    wb = openpyxl.load_workbook(BASE / "TV-Habitaciones.xlsx", data_only=True)
    ws = wb.active
    tvs = []
    for r in range(2, 36):
        hab = fmt(ws.cell(r, 1).value)
        marca = fmt(ws.cell(r, 2).value)
        if not hab or not marca:
            continue
        marca = marca.replace("otro: _________________", "").strip()
        pulg = fmt(ws.cell(r, 3).value).replace("otro: _______", "").strip()
        tvs.append({
            "zona": hab,
            "marca": marca[:40],
            "pulgadas": pulg[:20],
            "control_tv": fmt(ws.cell(r, 4).value) or "—",
            "control_deco": fmt(ws.cell(r, 5).value) or "—",
        })
    data["tvs"] = tvs
    wb.close()

    # JARDIN
    wb = openpyxl.load_workbook(BASE / "JArdineria" / "2024-Trabajos de jardinería 4 sectores .xlsx", data_only=True)
    jardin = []
    for sheet, sector in [("CENTRO", "Centro de eventos"), ("HOTEL1", "Hotel"), ("PARCELA CASA", "Casa / parcela")]:
        ws = wb[sheet]
        for r in range(2, (ws.max_row or 1) + 1):
            job = fmt(ws.cell(r, 2).value)
            if not job or len(job) < 4:
                continue
            prio = fmt(ws.cell(r, 1).value) or "—"
            inv = bool(ws.cell(r, 4).value)
            ver = bool(ws.cell(r, 6).value)
            temp = []
            if inv:
                temp.append("Invierno/otoño")
            if ver:
                temp.append("Verano/primavera")
            jardin.append({"sector": sector, "prioridad": prio, "faena": job[:140], "temporada": " · ".join(temp) or "—"})
    data["jardin"] = jardin
    data["riego"] = [
        {"dia": "Lunes", "zona": "Parcela casa"},
        {"dia": "Martes a viernes", "zona": "Parcela hotel"},
        {"dia": "Sábado y domingo", "zona": "Parcela casa"},
        {"dia": "Jueves 8:00–10:30", "zona": "Riego parcela casa (acequia)"},
        {"dia": "Lun / mié / vie", "zona": "Árboles Camino al Cielo"},
        {"dia": "Dom / lun / mié (noche)", "zona": "Piscina — 2 kg cloro"},
    ]
    data["labores"] = [
        {"sector": "Centro de eventos", "faena": "Orillar torre y bambú; Los Arados; sequía sala de máquinas; cosecha almendras/manzanas/ciruelas; poda crateus; corte bambú piscina chica"},
        {"sector": "Hotel", "faena": "Estrella, plato, hadas, riñón, mascanta norte, frontis y alrededor de recepción, plantas nuevas sector cerro este"},
        {"sector": "Permanente", "faena": "Riego Camino al Cielo L-M-V; árboles cancha y canil; plantas cocina C/E; manguera hotel; casa jefe orillar/pasto"},
    ]
    wb.close()

    data["prevencion"] = [
        {"id": "1.1", "tarea": "Reglamento interno", "quien": "Esteban Ortiz", "inicio": "01-01-2025", "fin": "31-01-2025", "estado": "Hecho"},
        {"id": "1.1.1", "tarea": "Protocolo TMERT", "quien": "Esteban Ortiz", "inicio": "10-02-2025", "fin": "10-03-2025", "estado": "Pendiente"},
        {"id": "1.2", "tarea": "Reunión psicosocial", "quien": "Esteban Ortiz", "inicio": "10-03-2025", "fin": "11-03-2025", "estado": "Pendiente"},
        {"id": "1.3", "tarea": "Difusión protocolo psicosocial", "quien": "Esteban Ortiz", "inicio": "12-03-2025", "fin": "13-03-2025", "estado": "Pendiente"},
        {"id": "1.4", "tarea": "Capacitación trabajo en altura", "quien": "Esteban Ortiz", "inicio": "12-03-2025", "fin": "31-03-2025", "estado": "Pendiente"},
    ]

    OUT.write_text("const MG = " + json.dumps(data, ensure_ascii=False, indent=2) + ";\n", encoding="utf-8")
    print("OK", OUT, {k: len(v) if isinstance(v, list) else v for k, v in data.items()})


if __name__ == "__main__":
    dump()
